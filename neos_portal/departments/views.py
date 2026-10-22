from django.contrib import messages
from django.db.models import Count, Prefetch, ProtectedError
from django.http import HttpResponse
from django.views.decorators.http import require_POST
from django.utils.http import content_disposition_header
from django.shortcuts import get_object_or_404, redirect, render

from . import stats as dept_stats
from .excel_utils import build_excel_columns

from .forms import (
    AssetCertForm,
    CertTypeForm,
    ClassificationEditForm,
    ClassificationForm,
    DepartmentForm,
    FieldDefinitionForm,
    asset_form_rows,
    build_asset_form,
)
from .models import (
    Asset,
    AssetCert,
    AssetFieldValue,
    CertType,
    Classification,
    ClassificationCertRequirement,
    Department,
    FieldDefinition,
)


def _parse_bulk_cert_filename(filename):
    """
    First underscore-separated part holds the expiry date, last part is the
    serial number. Anything in between (cert type, provider, description, in
    any combination) is ignored entirely, per the several filename formats in
    use. The date may be ddmmyyyy, or yyyymmdd behind an "EXP" marker -- see
    core/cert_filenames.py.
    """
    import os

    from core.cert_filenames import clean_serial, drop_marker_parts, read_date

    name_only = os.path.splitext(filename)[0]
    parts = drop_marker_parts(name_only.split("_"))
    if len(parts) < 2:
        return None, None, "Filename needs at least a date and a serial number separated by _."

    expiry_date = read_date(parts[0])
    if expiry_date is None:
        return None, None, (f"Couldn't read a date from '{parts[0].strip()}' (the first part of the filename "
                            "should be ddmmyyyy, or EXP followed by yyyymmdd).")
    return expiry_date, clean_serial(parts[-1]), None


def bulk_cert_upload(request):
    """Step 1: pick a cert type, upload files. Parses each filename and
    checks the serial against Departments, but saves nothing yet — the
    uploaded files are stashed on disk and a preview is shown next."""
    import uuid

    from django.conf import settings
    from django.core.files.storage import default_storage

    if request.method == "POST":
        cert_type_id = request.POST.get("cert_type")
        files = request.FILES.getlist("files")
        cert_type = CertType.objects.filter(pk=cert_type_id).first()

        if not cert_type:
            messages.error(request, "Choose a cert type first.")
            return redirect("departments:bulk_cert_upload")
        if not files:
            messages.error(request, "Choose at least one file to upload.")
            return redirect("departments:bulk_cert_upload")

        batch_id = uuid.uuid4().hex
        manifest = []
        for f in files:
            expiry_date, serial, error = _parse_bulk_cert_filename(f.name)
            entry = {"file": f.name, "expiry_date": None, "serial": serial, "error": error, "asset_id": None}

            if error:
                manifest.append(entry)
                continue

            asset = Asset.objects.filter(name=serial).first()
            if not asset:
                entry["error"] = f"No asset with Asset No/Serial No '{serial}' found in Departments."
                manifest.append(entry)
                continue

            temp_path = default_storage.save(f"tmp_bulk_certs/{batch_id}/{f.name}", f)
            entry["expiry_date"] = expiry_date.isoformat()
            entry["asset_id"] = asset.id
            entry["asset_name"] = asset.name
            entry["temp_path"] = temp_path
            manifest.append(entry)

        request.session[f"bulk_cert_batch_{batch_id}"] = {
            "cert_type_id": cert_type.id, "cert_type_name": cert_type.name, "manifest": manifest,
        }
        return redirect("departments:bulk_cert_preview", batch_id=batch_id)

    return render(request, "departments/bulk_cert_upload_results.html", {
        "cert_types": CertType.objects.all(),
    })


def bulk_cert_preview(request, batch_id):
    """Step 2: show what matched and what didn't, before anything is saved."""
    session_key = f"bulk_cert_batch_{batch_id}"
    batch = request.session.get(session_key)
    if not batch:
        messages.error(request, "That upload batch has expired or was already processed. Please upload again.")
        return redirect("departments:bulk_cert_upload")

    matched = [m for m in batch["manifest"] if not m["error"]]
    unmatched = [m for m in batch["manifest"] if m["error"]]

    return render(request, "departments/bulk_cert_preview.html", {
        "batch_id": batch_id, "cert_type_name": batch["cert_type_name"],
        "matched": matched, "unmatched": unmatched,
    })


def bulk_cert_confirm(request, batch_id):
    """Step 3: actually create/update the AssetCert records for the
    matched files, using the temp-saved file as the certificate file."""
    from django.core.files import File
    from django.core.files.storage import default_storage

    session_key = f"bulk_cert_batch_{batch_id}"
    batch = request.session.get(session_key)
    if not batch or request.method != "POST":
        messages.error(request, "Nothing to confirm — that batch has expired.")
        return redirect("departments:bulk_cert_upload")

    cert_type = CertType.objects.filter(pk=batch["cert_type_id"]).first()
    results = []

    for m in batch["manifest"]:
        if m["error"]:
            results.append({"file": m["file"], "ok": False, "reason": m["error"]})
            continue

        asset = Asset.objects.filter(pk=m["asset_id"]).first()
        if not asset or not cert_type:
            results.append({"file": m["file"], "ok": False, "reason": "Asset or cert type no longer exists."})
            continue

        import datetime
        expiry_date = datetime.date.fromisoformat(m["expiry_date"])

        with default_storage.open(m["temp_path"]) as temp_file:
            django_file = File(temp_file, name=m["file"])
            cert, created = AssetCert.objects.update_or_create(
                asset=asset, cert_type=cert_type,
                defaults={"expiry_date": expiry_date, "certificate_file": django_file},
            )
        default_storage.delete(m["temp_path"])

        results.append({
            "file": m["file"], "ok": True,
            "reason": f"{'Created' if created else 'Updated'} cert for {asset.name} \u2014 {cert_type.name}, expires {expiry_date}.",
        })

    del request.session[session_key]
    ok_count = sum(1 for r in results if r["ok"])
    messages.success(request, f"Saved {ok_count} of {len(results)} certificate(s).")

    return render(request, "departments/bulk_cert_upload_results.html", {
        "results": results, "cert_types": CertType.objects.all(),
    })


def department_list(request):
    departments = Department.objects.all()
    if request.method == "POST":
        form = DepartmentForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Department created.")
            return redirect("departments:list")
    else:
        form = DepartmentForm()
    departments = list(departments)
    per_department = dept_stats.department_stats(departments)
    for d in departments:
        d.stats = per_department[d.pk]
    return render(request, "departments/department_list.html", {
        "departments": departments, "form": form, "all_cert_types_global": CertType.objects.all(),
        "totals": dept_stats.total(per_department),
    })


def department_edit(request, pk):
    department = get_object_or_404(Department, pk=pk)
    if request.method == "POST":
        form = DepartmentForm(request.POST, instance=department)
        if form.is_valid():
            form.save()
            return redirect("departments:list")
    else:
        form = DepartmentForm(instance=department)
    return render(request, "departments/department_form.html", {"form": form, "department": department})


def department_delete(request, pk):
    department = get_object_or_404(Department, pk=pk)
    if request.method == "POST":
        try:
            department.delete()
            messages.success(request, "Department deleted.")
        except ProtectedError:
            messages.error(request, f"Can't delete '{department.name}' \u2014 job manifests still belong to it. Move or delete those first.")
        return redirect("departments:list")
    return render(request, "departments/department_confirm_delete.html", {"department": department})


def department_detail(request, pk):
    department = get_object_or_404(Department, pk=pk)
    tab = request.GET.get("tab", "assets")
    if request.method == "POST" and tab == "assets":
        form = ClassificationForm(request.POST, department=department)
        if form.is_valid():
            classification = form.save(commit=False)
            classification.department = department
            classification.save()
            source = form.cleaned_data.get("copy_fields_from")
            if source:
                for fd in source.fields.all():
                    FieldDefinition.objects.create(
                        classification=classification,
                        label=fd.label,
                        field_type=fd.field_type,
                        dropdown_options=fd.dropdown_options,
                        is_required=fd.is_required,
                        order=fd.order,
                    )
                messages.success(
                    request,
                    f"Classification created with {source.fields.count()} field(s) copied from {source.name}.",
                )
            return redirect(f"/departments/{department.pk}/?tab=assets")
    else:
        form = ClassificationForm(department=department)
    classifications = list(
        department.classifications.annotate(field_count=Count("fields", distinct=True)).order_by("order", "name")
    )
    per_classification = dept_stats.classification_stats(department=department)
    for c in classifications:
        c.stats = per_classification[c.pk]

    cert_classification_id = request.GET.get("classification", "")
    cert_type_id = request.GET.get("cert_type", "")
    if not cert_classification_id and classifications:
        # Default to showing just the FIRST classification's table instead
        # of every classification stacked one after another -- both for a
        # cleaner tabbed view and because rendering every classification's
        # assets/certs on every page load was the main source of the slow
        # loading time.
        cert_classification_id = str(classifications[0].pk)
    # Everything the certificates table needs is loaded up front, in a handful of
    # queries, instead of one query per asset / field / certificate cell.
    cert_classifications = list(
        Classification.objects.filter(pk=cert_classification_id, department=department).prefetch_related(
            "fields", "cert_requirements__cert_type",
            Prefetch("assets", queryset=Asset.objects.prefetch_related("values__field_definition", "certs")),
        )
    ) if cert_classification_id else []

    all_cert_types = CertType.objects.filter(
        classificationcertrequirement__classification__department=department
    ).distinct()

    # Certificates tab: asset fields are only shown as columns if picked in
    # the "Columns" chooser (default none, so it doesn't repeat the Assets tab).
    selected_cols = [int(c) for c in request.GET.getlist("cols") if c.isdigit()]
    cols_qs = "".join(f"&cols={c}" for c in selected_cols)

    return render(request, "departments/department_detail.html", {
        "selected_cols": selected_cols,
        "cols_qs": cols_qs,
        "department": department,
        "classifications": classifications,
        "totals": dept_stats.total(per_classification),
        "form": form,
        "tab": tab,
        "cert_classifications": cert_classifications,
        "cert_classification_id": cert_classification_id,
        "cert_type_id": cert_type_id,
        "all_cert_types": all_cert_types,
        "all_cert_types_global": CertType.objects.all(),
    })


def classification_edit(request, pk):
    classification = get_object_or_404(Classification, pk=pk)
    form = ClassificationEditForm(request.POST or None, instance=classification)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, f"'{classification.name}' updated.")
        return redirect(f"/departments/{classification.department_id}/?tab=assets")
    return render(request, "departments/classification_form.html", {"form": form, "classification": classification})


@require_POST
def classification_delete(request, pk):
    classification = get_object_or_404(Classification, pk=pk)
    department_pk = classification.department_id
    classification.delete()
    return redirect(f"/departments/{department_pk}/?tab=assets")


def classification_detail(request, pk):
    """Assets view for a classification. Field configuration is a popup on
    this same page (a New-field form + field list inside a <dialog>),
    rather than a separate page navigation."""
    classification = get_object_or_404(Classification, pk=pk)

    if request.method == "POST":
        field_form = FieldDefinitionForm(request.POST)
        if field_form.is_valid():
            fd = field_form.save(commit=False)
            fd.classification = classification
            fd.save()
            return redirect(f"/departments/classification/{pk}/?open_fields=1")
    else:
        field_form = FieldDefinitionForm()

    field_defs = classification.fields.all()
    assets = classification.assets.prefetch_related("values__field_definition", "approvals")
    asset_rows = [{"asset": a, "values": a.values_by_field_id()} for a in assets]
    stats = dept_stats.classification_stats(classification=classification)[classification.pk]
    stats["at_workshop"] = stats["assets"] - stats["away"]

    # For classifications of units, show the NPT logged against each unit.
    npt_month_total = npt_all_total = None
    if classification.is_unit:
        import datetime

        from django.db.models import Q, Sum

        from sqb.models import SQBTrackerEntry

        first = datetime.date.today().replace(day=1)
        sums = {
            r["unit"]: r for r in SQBTrackerEntry.objects.filter(unit__classification=classification)
            .values("unit").annotate(total=Sum("npt_hours"), month=Sum("npt_hours", filter=Q(date__gte=first)))
        }
        for row in asset_rows:
            found = sums.get(row["asset"].pk, {})
            row["npt_month"] = round(found.get("month") or 0, 1)
            row["npt_total"] = round(found.get("total") or 0, 1)
        npt_month_total = round(sum(r["npt_month"] for r in asset_rows), 1)
        npt_all_total = round(sum(r["npt_total"] for r in asset_rows), 1)

    return render(request, "departments/classification_detail.html", {
        "classification": classification,
        "stats": stats,
        "npt_month_total": npt_month_total, "npt_all_total": npt_all_total,
        "field_defs": field_defs,
        "asset_rows": asset_rows,
        "field_form": field_form,
        "open_fields": request.GET.get("open_fields") or bool(request.method == "POST" and field_form.errors),
    })


def classification_fields(request, pk):
    """Field-definition management for a classification \u2014 separate page
    from the assets list, reached via a 'Configure fields' link."""
    classification = get_object_or_404(Classification, pk=pk)
    field_defs = classification.fields.all()

    if request.method == "POST":
        form = FieldDefinitionForm(request.POST)
        if form.is_valid():
            fd = form.save(commit=False)
            fd.classification = classification
            fd.save()
            return redirect("departments:classification_fields", pk=pk)
    else:
        form = FieldDefinitionForm()

    return render(request, "departments/classification_fields.html", {
        "classification": classification,
        "field_defs": field_defs,
        "form": form,
    })


@require_POST
def field_delete(request, pk):
    fd = get_object_or_404(FieldDefinition, pk=pk)
    classification_pk = fd.classification_id
    fd.delete()
    return redirect(f"/departments/classification/{classification_pk}/?open_fields=1")


def classification_export_xlsx(request, pk):
    import openpyxl

    classification = get_object_or_404(Classification, pk=pk)
    field_defs = list(classification.fields.all())
    headers, col_specs = build_excel_columns(field_defs)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = classification.name[:31]
    ws.append(headers)

    for asset in classification.assets.all():
        raw = {v.field_definition_id: v for v in asset.values.all()}
        row = [asset.name, asset.description]
        for kind, fd in col_specs[2:]:
            v = raw.get(fd.id)
            if kind == "value_unit_value":
                row.append(v.number_value if v else None)
            elif kind == "value_unit_unit":
                row.append(v.unit_value if v else "")
            elif kind == "attachment":
                row.append(str(v.file_value) if v and v.file_value else "")
            elif fd.field_type == FieldDefinition.NUMBER:
                row.append(v.number_value if v else None)
            else:
                row.append(v.text_value if v else "")
        ws.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = content_disposition_header(True, f"{classification.name}_assets.xlsx")
    wb.save(response)
    return response


def classification_export_selected_xlsx(request, pk):
    import openpyxl

    classification = get_object_or_404(Classification, pk=pk)
    if request.method != "POST":
        return redirect("departments:classification_detail", pk=pk)

    asset_ids = request.POST.getlist("asset_ids")
    field_ids = {int(i) for i in request.POST.getlist("field_ids")}

    field_defs = [fd for fd in classification.fields.all() if fd.id in field_ids]
    headers, col_specs = build_excel_columns(field_defs)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = classification.name[:31]
    ws.append(headers)

    assets = classification.assets.filter(id__in=asset_ids) if asset_ids else classification.assets.none()
    for asset in assets:
        raw = {v.field_definition_id: v for v in asset.values.all()}
        row = [asset.name, asset.description]
        for kind, fd in col_specs[2:]:
            v = raw.get(fd.id)
            if kind == "value_unit_value":
                row.append(v.number_value if v else None)
            elif kind == "value_unit_unit":
                row.append(v.unit_value if v else "")
            elif kind == "attachment":
                row.append(str(v.file_value) if v and v.file_value else "")
            elif fd.field_type == FieldDefinition.NUMBER:
                row.append(v.number_value if v else None)
            else:
                row.append(v.text_value if v else "")
        ws.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = content_disposition_header(True, f"{classification.name}_selected.xlsx")
    wb.save(response)
    return response


def classification_template_xlsx(request, pk):
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    classification = get_object_or_404(Classification, pk=pk)
    field_defs = list(classification.fields.all())
    headers, col_specs = build_excel_columns(field_defs)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"
    ws.append(headers)

    for idx, (kind, fd) in enumerate(col_specs, start=1):
        if fd is None:
            continue
        needs_dropdown = (
            (kind == "simple" and fd.field_type in (FieldDefinition.DROPDOWN, FieldDefinition.SEARCHABLE))
            or kind == "value_unit_unit"
        )
        if needs_dropdown:
            options = fd.options_list()
            if options:
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                col_letter = get_column_letter(idx)
                dv.add(f"{col_letter}2:{col_letter}1000")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = content_disposition_header(True, f"{classification.name}_template.xlsx")
    wb.save(response)
    return response


def classification_import_xlsx(request, pk):
    import openpyxl

    classification = get_object_or_404(Classification, pk=pk)
    field_defs = list(classification.fields.all())
    headers, col_specs = build_excel_columns(field_defs)

    if request.method == "POST" and request.FILES.get("file"):
        wb = openpyxl.load_workbook(request.FILES["file"], data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            messages.error(request, "That file has no rows.")
            return redirect("departments:classification_import", pk=pk)

        header_row = [str(h).strip() if h is not None else "" for h in rows[0]]
        if header_row != headers:
            messages.warning(
                request,
                "Column headers didn't exactly match the current template \u2014 "
                "matched by position instead. Re-download the template if you've "
                "added or removed fields since your last export.",
            )

        created = 0
        skipped = 0
        duplicates = 0
        missing_description = 0
        for data_row in rows[1:]:
            if not data_row or data_row[0] in (None, ""):
                skipped += 1
                continue
            name = str(data_row[0]).strip()
            description = str(data_row[1]).strip() if len(data_row) > 1 and data_row[1] not in (None, "") else ""
            if not description:
                missing_description += 1
                continue
            if Asset.objects.filter(name=name).exists():
                duplicates += 1
                continue
            asset = Asset.objects.create(classification=classification, name=name, description=description)

            col_idx = 2
            for kind, fd in col_specs[2:]:
                value = data_row[col_idx] if col_idx < len(data_row) else None
                if kind == "attachment":
                    col_idx += 1
                    continue
                afv, _ = AssetFieldValue.objects.get_or_create(asset=asset, field_definition=fd)
                if kind == "value_unit_value":
                    try:
                        afv.number_value = float(value) if value not in (None, "") else None
                    except (TypeError, ValueError):
                        afv.number_value = None
                elif kind == "value_unit_unit":
                    afv.unit_value = str(value) if value else ""
                elif fd.field_type == FieldDefinition.NUMBER:
                    try:
                        afv.number_value = float(value) if value not in (None, "") else None
                    except (TypeError, ValueError):
                        afv.number_value = None
                else:
                    afv.text_value = str(value) if value is not None else ""
                afv.save()
                col_idx += 1
            created += 1

        messages.success(
            request,
            f"Imported {created} asset(s). Skipped {skipped} blank row(s), "
            f"{missing_description} row(s) missing a required Description, "
            f"{duplicates} duplicate Asset No/Serial No already in the system.",
        )
        return redirect("departments:classification_detail", pk=pk)

    return render(request, "departments/classification_import.html", {"classification": classification})


def _save_field_value(asset, fd, form):
    """Writes one AssetFieldValue from cleaned form data, handling the
    Value-with-unit field's two sub-inputs as a special case."""
    afv, _ = AssetFieldValue.objects.get_or_create(asset=asset, field_definition=fd)
    if fd.field_type == FieldDefinition.NUMBER:
        afv.number_value = form.cleaned_data.get(fd.field_name())
    elif fd.field_type == FieldDefinition.VALUE_UNIT:
        afv.number_value = form.cleaned_data.get(f"{fd.field_name()}_value")
        afv.unit_value = form.cleaned_data.get(f"{fd.field_name()}_unit") or ""
    elif fd.field_type == FieldDefinition.ATTACHMENT:
        value = form.cleaned_data.get(fd.field_name())
        if value:
            afv.file_value = value
    else:
        afv.text_value = form.cleaned_data.get(fd.field_name()) or ""
    afv.save()


def asset_create(request, classification_pk):
    classification = get_object_or_404(Classification, pk=classification_pk)
    FormClass, field_defs = build_asset_form(classification)
    if request.method == "POST":
        form = FormClass(request.POST, request.FILES)
        if form.is_valid():
            asset = Asset.objects.create(
                classification=classification,
                name=form.cleaned_data["name"],
                description=form.cleaned_data["description"],
            )
            for fd in field_defs:
                _save_field_value(asset, fd, form)
            messages.success(request, "Asset added.")
            return redirect("departments:classification_detail", pk=classification_pk)
    else:
        form = FormClass()
    rows = asset_form_rows(form, field_defs)
    return render(request, "departments/asset_form.html", {
        "form": form, "rows": rows, "classification": classification,
    })


def asset_detail(request, pk):
    """Everything about one piece of equipment in one place: certificates, maintenance jobs, QHSE reports and the
    open actions that came out of them. Each of these is still kept in, and edited from, its own tab."""
    from action_tracker.models import Action
    from maintenance.models import MaintenanceRecord
    from qhse.models import Incident, Inspection

    asset = get_object_or_404(Asset.objects.select_related("classification__department"), pk=pk)
    jobs = list(MaintenanceRecord.objects.filter(asset_serial=asset.name).prefetch_related("actions"))
    incidents = list(Incident.objects.filter(asset_serial=asset.name).prefetch_related("actions"))
    inspections = list(Inspection.objects.filter(asset_serial=asset.name).prefetch_related("actions"))
    linked = {}
    for source in jobs + incidents + inspections:
        for a in source.actions.all():
            linked[a.pk] = a
    actions = sorted(linked.values(), key=Action.urgency_key)
    return render(request, "departments/asset_detail.html", {
        "asset": asset, "certs": asset.certs.select_related("cert_type"),
        "jobs": jobs, "incidents": incidents, "inspections": inspections, "actions": actions,
        "open_actions": sum(a.is_open for a in actions),
        "overdue_jobs": sum(j.status == "expired" for j in jobs),
    })


def asset_edit(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    classification = asset.classification
    FormClass, field_defs = build_asset_form(classification)

    if request.method == "POST":
        form = FormClass(request.POST, request.FILES, exclude_asset_id=asset.id)
        if form.is_valid():
            asset.name = form.cleaned_data["name"]
            asset.description = form.cleaned_data["description"]
            asset.save()
            for fd in field_defs:
                _save_field_value(asset, fd, form)
            return redirect("departments:classification_detail", pk=classification.pk)
    else:
        initial = {"name": asset.name, "description": asset.description}
        existing = {v.field_definition_id: v for v in asset.values.all()}
        for fd in field_defs:
            v = existing.get(fd.id)
            if not v:
                continue
            if fd.field_type == FieldDefinition.NUMBER:
                initial[fd.field_name()] = v.number_value
            elif fd.field_type == FieldDefinition.VALUE_UNIT:
                initial[f"{fd.field_name()}_value"] = v.number_value
                initial[f"{fd.field_name()}_unit"] = v.unit_value
            elif fd.field_type != FieldDefinition.ATTACHMENT:
                initial[fd.field_name()] = v.text_value
        form = FormClass(initial=initial, exclude_asset_id=asset.id)

    rows = asset_form_rows(form, field_defs)
    return render(request, "departments/asset_form.html", {
        "form": form, "rows": rows, "classification": classification, "asset": asset,
    })


@require_POST
def asset_delete(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    classification_pk = asset.classification_id
    asset.delete()
    return redirect("departments:classification_detail", pk=classification_pk)


@require_POST
def asset_request_removal(request, pk):
    """Level 2 can't delete an asset; it asks, and a Level 3 approves (which deletes it)."""
    from approvals import services as approvals

    asset = get_object_or_404(Asset, pk=pk)
    if request.user.profile.can_delete:
        asset.delete()
        messages.success(request, f"{asset.name} removed.")
    else:
        approvals.submit(asset, "asset_removal", request.user)
        messages.success(request, f"Removal of {asset.name} sent to a Level 3 for approval. It stays until then.")
    return redirect("departments:classification_detail", pk=asset.classification_id)


def cert_type_manage(request):
    cert_types = CertType.objects.all()
    if request.method == "POST":
        form = CertTypeForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect("departments:cert_type_manage")
    else:
        form = CertTypeForm()
    return render(request, "departments/cert_type_manage.html", {"cert_types": cert_types, "form": form})


@require_POST
def cert_type_delete(request, pk):
    ct = get_object_or_404(CertType, pk=pk)
    ct.delete()
    return redirect("departments:cert_type_manage")


def classification_certtypes(request, pk):
    classification = get_object_or_404(Classification, pk=pk)
    all_cert_types = CertType.objects.all()
    assigned_ids = set(classification.cert_requirements.values_list("cert_type_id", flat=True))

    if request.method == "POST":
        selected_ids = {int(i) for i in request.POST.getlist("cert_types")}
        ClassificationCertRequirement.objects.filter(classification=classification).exclude(
            cert_type_id__in=selected_ids
        ).delete()
        for cid in selected_ids - assigned_ids:
            ClassificationCertRequirement.objects.create(classification=classification, cert_type_id=cid)
        return redirect(f"/departments/{classification.department_id}/?tab=certificates")

    return render(request, "departments/classification_certtypes.html", {
        "classification": classification,
        "all_cert_types": all_cert_types,
        "assigned_ids": assigned_ids,
    })


def asset_cert_create(request, asset_pk):
    asset = get_object_or_404(Asset, pk=asset_pk)
    classification = asset.classification
    if request.method == "POST":
        form = AssetCertForm(request.POST, request.FILES, classification=classification)
        if form.is_valid():
            cert = form.save(commit=False)
            cert.asset = asset
            cert.save()
            return redirect(f"/departments/{classification.department_id}/?tab=certificates")
    else:
        form = AssetCertForm(classification=classification)
    return render(request, "departments/asset_cert_form.html", {"form": form, "asset": asset})


def asset_cert_edit(request, pk):
    cert = get_object_or_404(AssetCert, pk=pk)
    classification = cert.asset.classification
    if request.method == "POST":
        form = AssetCertForm(request.POST, request.FILES, instance=cert, classification=classification)
        if form.is_valid():
            form.save()
            return redirect(f"/departments/{classification.department_id}/?tab=certificates")
    else:
        form = AssetCertForm(instance=cert, classification=classification)
    return render(request, "departments/asset_cert_form.html", {"form": form, "asset": cert.asset})
