import io

from django.contrib import messages
from django.db.models import ProtectedError, Q
from django.http import HttpResponse
from django.utils.http import content_disposition_header
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from departments.models import Asset

from approvals import services as approvals

from .forms import JobManifestForm
from .models import EquipmentTransfer, JobManifest
from .transfers import TransferError, transfer_serials


def _hold_back_if_unapproved(request, job):
    """A manifest can't go Active (dispatch the equipment) until a Level 3 has approved it."""
    if job.status == "Active" and not approvals.is_approved(job, "manifest"):
        JobManifest.objects.filter(pk=job.pk).update(status="Draft")
        job.status = "Draft"
        messages.warning(request, "This manifest is waiting for approval, so it stays a Draft. "
                                  "Once a Level 3 approves it, set it to Active to dispatch the equipment.")
        return True
    return False


def job_list(request):
    jobs = JobManifest.objects.select_related("department").prefetch_related("approvals").all()
    q = request.GET.get("q")
    if q:
        jobs = jobs.filter(job_name__icontains=q) | jobs.filter(client__icontains=q)
    return render(request, "job_manager/job_list.html", {"jobs": jobs, "q": q or ""})


def job_create(request):
    if request.method == "POST":
        form = JobManifestForm(request.POST)
        if form.is_valid():
            job = form.save()
            messages.success(request, f"Job '{job.job_name}' created ({job.job_number}).")
            _hold_back_if_unapproved(request, job)
            if job.status == "Active":
                job.dispatch_equipment()
            return redirect("job_manager:detail", pk=job.pk)
    else:
        form = JobManifestForm()
    return render(request, "job_manager/job_form.html", {"form": form})


def job_edit(request, pk):
    job = get_object_or_404(JobManifest, pk=pk)
    old_status = job.status
    if request.method == "POST":
        form = JobManifestForm(request.POST, instance=job)
        if form.is_valid():
            job = form.save()
            if old_status != "Active" and job.status == "Active" and _hold_back_if_unapproved(request, job):
                return redirect("job_manager:detail", pk=pk)
            if old_status != "Active" and job.status == "Active":
                job.dispatch_equipment()
                messages.success(request, f"Equipment dispatched to '{job.location or 'the job site'}'.")
            elif old_status != "Completed" and job.status == "Completed":
                job.return_equipment()
                messages.success(request, "All equipment on this manifest returned to Workshop.")
            return redirect("job_manager:detail", pk=pk)
    else:
        form = JobManifestForm(instance=job)
    return render(request, "job_manager/job_form.html", {"form": form, "job": job})


def job_delete(request, pk):
    job = get_object_or_404(JobManifest, pk=pk)
    if request.method == "POST":
        try:
            job.delete()
            messages.success(request, "Job deleted.")
        except ProtectedError:
            report_names = ", ".join(r.name for r in job.job_reports.all())
            messages.error(
                request,
                f"Can't delete this manifest \u2014 it still has job report(s) linked to it ({report_names}). "
                "Delete or reassign those job reports first.",
            )
        return redirect("job_manager:list")
    return render(request, "job_manager/job_confirm_delete.html", {"job": job})


def job_manifest_import_serials(request, pk):
    """Lets a supervisor upload the Excel sheet where they've already
    written down serial numbers by hand, and shows which ones actually
    exist in the database (green) vs don't (red) before adding the
    matched ones. Reuses job_detail's existing add_serials handling --
    the preview's checkboxes post straight to that view with the same
    add_serial field name, so there's no separate "confirm" step or
    server-side state to carry between requests."""
    import openpyxl

    job = get_object_or_404(JobManifest, pk=pk)
    rows = None
    if request.method == "POST" and request.FILES.get("excel_file"):
        wb = openpyxl.load_workbook(request.FILES["excel_file"], data_only=True)
        ws = wb.active
        existing = set(job.serial_list())
        taken = JobManifest.reserved_map(exclude_pk=job.pk)
        header_words = {"serial", "serial number", "asset no", "asset no / serial no", "asset no/serial no"}
        rows = []
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] in (None, ""):
                continue
            serial = str(row[0]).strip()
            if serial.lower() in header_words:
                continue
            if serial in existing:
                continue
            asset = Asset.objects.filter(name=serial).select_related(
                "classification", "classification__department"
            ).first()
            rows.append({
                "serial": serial, "asset": asset, "found": asset is not None,
                "taken_by": taken.get(serial),
                "away": asset is not None and asset.current_location != "Workshop" and serial not in taken,
            })

    return render(request, "job_manager/job_import_serials.html", {"job": job, "rows": rows})


def job_detail(request, pk):
    job = get_object_or_404(JobManifest, pk=pk)

    if request.method == "POST":
        if "add_serials" in request.POST:
            requested = request.POST.getlist("add_serial")
            taken = JobManifest.reserved_map(exclude_pk=job.pk)
            unapproved = approvals.blocked_asset_names(requested)
            new_serials = []
            for serial in requested:
                if serial in unapproved:
                    messages.error(request, f"{serial} is waiting for approval (or a removal request), so it can't go on a manifest yet.")
                elif serial in taken:
                    messages.error(request, f"{serial} is already on manifest '{taken[serial].job_name}' "
                                            f"({taken[serial].job_number}) \u2014 transfer it from there instead.")
                else:
                    new_serials.append(serial)
            current = job.serial_list()
            added = 0
            for s in new_serials:
                if s not in current:
                    current.append(s)
                    added += 1
            job.equipment_serials = ",".join(current)
            job.save()
            # If the manifest is already Active, anything added now should
            # be dispatched immediately rather than waiting for a future
            # status change (which has already happened for this manifest).
            if job.status == "Active" and new_serials:
                Asset.objects.filter(name__in=new_serials).update(
                    current_location=job.location or "Dispatched"
                )
                messages.success(request, f"Added {added} item(s) and dispatched them (manifest is Active).")
            else:
                messages.success(request, f"Added {added} item(s) to the manifest.")
            return redirect(f"{reverse('job_manager:detail', args=[pk])}?tab=add")
        elif "remove_serial" in request.POST:
            serial = request.POST["remove_serial"]
            current = [s for s in job.serial_list() if s != serial]
            job.equipment_serials = ",".join(current)
            job.save()
        elif "transfer_serials" in request.POST:
            chosen = request.POST.getlist("transfer_serial")
            target = request.POST.get("destination", "")
            note = request.POST.get("note", "").strip()
            if not chosen:
                messages.error(request, "Tick the items you want to transfer first.")
            elif not target:
                messages.error(request, "Choose where to transfer them to.")
            else:
                destination = None if target == "workshop" else get_object_or_404(JobManifest, pk=target)
                try:
                    moved = transfer_serials(chosen, job, destination, user=request.user, note=note)
                    where = "the Workshop" if destination is None else f"'{destination.job_name}'"
                    messages.success(request, f"Transferred {len(moved)} item(s) to {where}.")
                except TransferError as exc:
                    messages.error(request, str(exc))
        elif "return_serial" in request.POST:
            serial = request.POST["return_serial"]
            Asset.objects.filter(name=serial).update(current_location="Workshop")
            messages.success(request, f"'{serial}' returned to Workshop and is available for other jobs again.")
        return redirect("job_manager:detail", pk=pk)

    existing = set(job.serial_list())
    taken = JobManifest.reserved_map(exclude_pk=job.pk)  # items already on another open manifest

    # Free-text search (kept alongside the classification browser below),
    # scoped to the manifest's own department once one is set.
    search = request.GET.get("search", "")
    search_results = []
    if search:
        qs = Asset.objects.filter(name__icontains=search, current_location="Workshop")
        if job.department_id:
            qs = qs.filter(classification__department=job.department)
        search_results = [
            a for a in qs.select_related("classification", "classification__department")[:30]
            if a.name not in existing and a.name not in taken
        ]

    # Browse-by-classification: every classification under this manifest's
    # department, each with its own list of Workshop-available assets not
    # already on this manifest -- lets you see the whole department's
    # equipment organized the way it's actually organized, instead of only
    # searching by serial number.
    classification_groups = []
    if job.department_id:
        for classification in job.department.classifications.all():
            available = [
                a for a in classification.assets.filter(current_location="Workshop").prefetch_related("values")
                if a.name not in existing and a.name not in taken
            ]
            field_defs = list(classification.fields.all())
            rows = [{"asset": a, "values": a.values_by_field_id()} for a in available]
            classification_groups.append({
                "classification": classification,
                "field_defs": field_defs,
                "rows": rows,
                "count": len(available),
            })

    manifest_rows = []
    for a in job.assets():
        if not a.is_available:
            state = "Dispatched"
        elif job.status == "Draft":
            state = "Reserved (draft)"
        else:
            state = "At Workshop"
        manifest_rows.append({"asset": a, "state": state, "also_on": taken.get(a.name),
                              "where": a.current_location if not a.is_available else ""})

    from collections import Counter
    by_class = sorted(Counter(r["asset"].classification.name for r in manifest_rows).items())
    stats = {
        "total": len(manifest_rows),
        "dispatched": sum(1 for r in manifest_rows if r["state"] == "Dispatched"),
        "reserved": sum(1 for r in manifest_rows if r["state"] == "Reserved (draft)"),
        "at_workshop": sum(1 for r in manifest_rows if r["state"] == "At Workshop"),
        "conflicts": sum(1 for r in manifest_rows if r["also_on"]),
    }
    tab = request.GET.get("tab") or ("add" if search else "manifest")
    if tab not in ("manifest", "add", "history"):
        tab = "manifest"

    other_manifests = JobManifest.objects.exclude(status="Completed").exclude(pk=job.pk).order_by("-date", "job_name")

    return render(request, "job_manager/job_detail.html", {
        "job": job,
        "assets": job.assets(),
        "manifest_rows": manifest_rows,
        "stats": stats, "by_class": by_class, "tab": tab,
        "other_manifests": other_manifests,
        "transfers": EquipmentTransfer.objects.filter(
            Q(from_manifest=job) | Q(to_manifest=job)).select_related("moved_by")[:15],
        "search": search,
        "search_results": search_results,
        "classification_groups": classification_groups,
    })


def job_export_manifest_pdf(request, pk):
    from .manifest_pdf import build_manifest_pdf

    job = get_object_or_404(JobManifest, pk=pk)
    buffer = build_manifest_pdf(job)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = content_disposition_header(True, f"Manifest_{job.job_number}.pdf")
    return response


def job_export_manifest_xlsx(request, pk):
    import openpyxl

    job = get_object_or_404(JobManifest, pk=pk)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    assets_by_classification = {}
    for asset in job.assets():
        assets_by_classification.setdefault(asset.classification, []).append(asset)

    if not assets_by_classification:
        ws = wb.create_sheet("Manifest")
        ws.append(["No equipment on this manifest yet."])
    else:
        for classification, assets in assets_by_classification.items():
            # Sheet names can't exceed 31 chars or contain []:*?/\
            safe_name = "".join(c for c in classification.name if c not in "[]:*?/\\")[:31] or "Sheet"
            ws = wb.create_sheet(safe_name)
            # Location isn't part of the manifest itself -- it only tracks
            # where equipment lives in the warehouse -- so it's dropped in
            # favor of the classification's own fields (matching the PDF
            # export, which already shows these).
            field_defs = list(classification.fields.all())
            ws.append(["Asset No / Serial No", "Description"] + [fd.label for fd in field_defs])
            for asset in assets:
                values = asset.values_by_field_id()
                row = [asset.name, asset.description]
                for fd in field_defs:
                    v = values.get(fd.id)
                    row.append(v if v not in (None, "") else "\u2014")
                ws.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = content_disposition_header(True, f"Manifest_{job.job_number}.xlsx")
    wb.save(response)
    return response


def job_export_certs_zip(request, pk):
    """Bundles every uploaded certificate FILE (not just the status
    table) into a single ZIP, organized into one folder per
    classification, so you can hand over "all the certs for the
    selected instruments in one folder" per section."""
    import zipfile

    job = get_object_or_404(JobManifest, pk=pk)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        any_files = False
        for asset in job.assets():
            folder = asset.classification.name.replace("/", "-").strip() or "Uncategorized"
            for cert in asset.certs.select_related("cert_type").all():
                if not cert.certificate_file:
                    continue
                try:
                    cert.certificate_file.open("rb")
                    data = cert.certificate_file.read()
                    cert.certificate_file.close()
                except Exception:
                    continue
                original_name = cert.certificate_file.name
                ext = original_name.rsplit(".", 1)[-1] if "." in original_name else "pdf"
                safe_asset = asset.name.replace("/", "-")
                safe_cert_type = cert.cert_type.name.replace("/", "-")
                arcname = f"{folder}/{safe_asset}_{safe_cert_type}.{ext}".replace(" ", "_")
                zf.writestr(arcname, data)
                any_files = True
        if not any_files:
            zf.writestr("README.txt", "No certificate files have been uploaded for any asset on this manifest yet.")

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = content_disposition_header(True, f"Certs_{job.job_number}.zip")
    return response


def job_export_certs_xlsx(request, pk):
    import openpyxl

    job = get_object_or_404(JobManifest, pk=pk)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Certificates"
    ws.append(["Asset No / Serial No", "Cert Type", "Status", "Expiry Date"])
    for asset in job.assets():
        for cert in asset.certs.select_related("cert_type").all():
            ws.append([asset.name, cert.cert_type.name, cert.status, str(cert.expiry_date or "")])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = content_disposition_header(True, f"Certs_{job.job_number}.xlsx")
    wb.save(response)
    return response